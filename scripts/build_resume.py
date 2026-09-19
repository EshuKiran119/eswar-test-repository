from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
import json, argparse

BASE=Path(__file__).resolve().parents[1]
parser=argparse.ArgumentParser()
parser.add_argument('--verified-portfolio', required=True, help='Public portfolio URL after deployment verification')
parser.add_argument('--output-dir', default=str(BASE/'build/resume'))
args=parser.parse_args()
# User-specified resume visual system: A4, single column, Arial, navy hierarchy.
# Every value below is explicit; no tables, text boxes, or fake list paragraphs.
doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.2677); sec.page_height=Inches(11.6929)
sec.top_margin=sec.bottom_margin=Inches(.62)
sec.left_margin=sec.right_margin=Inches(.65)
sec.header_distance=sec.footer_distance=Inches(.3)
for name,size,bold,color,before,after in [
 ('Normal',10.5,False,'202B39',0,4),
 ('Title',20,True,'132D46',0,4),
 ('Subtitle',11.3,True,'244D6B',0,5),
 ('Heading 1',10.5,True,'244D6B',11,5),
 ('Heading 2',11,True,'132D46',4,2),
 ('Heading 3',10.5,True,'132D46',4,2),
 ('List Bullet',10.5,False,'202B39',0,7),
 ('Contact',9.3,False,'344456',0,3),
 ('Role Meta',9.6,False,'344456',0,7),
 ('Skill',10,False,'202B39',0,5),
 ('Footer',8,False,'526272',0,0)]:
 s=doc.styles[name] if name in doc.styles else doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
 s.font.italic=False
 s.font.name='Arial';s.font.size=Pt(size);s.font.bold=bold;s.font.color.rgb=RGBColor.from_string(color)
 s.paragraph_format.space_before=Pt(before);s.paragraph_format.space_after=Pt(after)
 s.paragraph_format.line_spacing=1.08
 s.paragraph_format.widow_control=True
 s.paragraph_format.keep_with_next=name.startswith('Heading') or name in ['Title','Subtitle','Contact','Role Meta']
 s.paragraph_format.keep_together=True
 s.paragraph_format.left_indent=Inches(0);s.paragraph_format.right_indent=Inches(0)
 s.paragraph_format.first_line_indent=Inches(0)
# Remove inherited template decoration and character tracking.
for st in doc.styles:
 for el in list(st.element.iter()):
  if el.tag in [qn('w:pBdr'),qn('w:spacing')] and el.getparent().tag==qn('w:rPr') or el.tag in [qn('w:pBdr'),qn('w:contextualSpacing')]:
   el.getparent().remove(el)
numroot=doc.part.numbering_part.element
abstract=OxmlElement('w:abstractNum');abstract.set(qn('w:abstractNumId'),'42')
lvl=OxmlElement('w:lvl');lvl.set(qn('w:ilvl'),'0')
for tag,val in [('start','1'),('numFmt','bullet'),('lvlText','•'),('lvlJc','left')]:
 e=OxmlElement('w:'+tag);e.set(qn('w:val'),val);lvl.append(e)
ppr=OxmlElement('w:pPr'); ind=OxmlElement('w:ind');ind.set(qn('w:left'),'220');ind.set(qn('w:hanging'),'220');ppr.append(ind)
tabs=OxmlElement('w:tabs');tab=OxmlElement('w:tab');tab.set(qn('w:val'),'num');tab.set(qn('w:pos'),'220');tabs.append(tab);ppr.append(tabs);lvl.append(ppr)
abstract.append(lvl);numroot.append(abstract)
num=OxmlElement('w:num');num.set(qn('w:numId'),'42');e=OxmlElement('w:abstractNumId');e.set(qn('w:val'),'42');num.append(e);numroot.append(num)

def p(text='',style=None):return doc.add_paragraph(text,style)
def bullet(text):
 para=p(text,'List Bullet');pr=para._p.get_or_add_pPr();n=OxmlElement('w:numPr')
 for tag,val in [('ilvl','0'),('numId','42')]:
  e=OxmlElement('w:'+tag);e.set(qn('w:val'),val);n.append(e)
 pr.append(n);para.paragraph_format.left_indent=Inches(220/1440);para.paragraph_format.first_line_indent=Inches(-220/1440)
def link(para,label,url):
 rid=para.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
 h=OxmlElement('w:hyperlink');h.set(qn('r:id'),rid);r=OxmlElement('w:r');pr=OxmlElement('w:rPr')
 c=OxmlElement('w:color');c.set(qn('w:val'),'244D6B');pr.append(c)
 u=OxmlElement('w:u');u.set(qn('w:val'),'single');pr.append(u);r.append(pr)
 t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);para._p.append(h)
def heading(text):p(text,'Heading 1')
def role(title,company,dates):
 p(title,'Heading 2');p(company+' | Hyderabad, India | '+dates,'Role Meta')
def labelled(label,text,style='Skill'):
 para=p('',style);para.add_run(label+': ').bold=True;para.add_run(text);return para

NAME='ESWAR SAI KIRAN SINGAMSETTY'
GH='https://github.com/EshuKiran119/eswar-test-repository'
LI='https://www.linkedin.com/in/eswar-sai-kiran-singamsetty-429840198'
p(NAME,'Title');p('Senior QA Engineer | SDET | Automation & Quality Engineering','Subtitle')
c=p('Hyderabad, India | +91 63014 49624 | ','Contact');link(c,'eshukiran57@gmail.com','mailto:eshukiran57@gmail.com')
c=p('','Contact');link(c,'LinkedIn',LI);c.add_run(' | ');link(c,'GitHub: EshuKiran119/eswar-test-repository',GH)
if args.verified_portfolio:
 c=p('Portfolio: ','Contact');link(c,args.verified_portfolio.removeprefix('https://'),args.verified_portfolio)
heading('PROFESSIONAL SUMMARY')
p('Senior QA Engineer and SDET with nearly 5 years of experience across UI, API, performance and security testing. Designed and owns the Playwright, Python and pytest-bdd framework at Experian, integrating UI/API automation with CI/CD release gates and AWS observability. Combines specification-driven quality gates, cross-team QA ownership and practical AI-assisted testing to improve testability, automation readiness and release confidence.')
heading('CORE ENGINEERING EXPERTISE')
p('Automation framework architecture | UI and API automation | Authentication and data validation | Performance engineering | CI/CD quality gates | Cloud troubleshooting | Shift-left QA','Skill')
heading('PROFESSIONAL EXPERIENCE')
role('Senior QA Engineer','Experian Services India Private Limited','Jul 2025 - Present')
exp=[
'Designed and built the Playwright + Python + pytest-bdd automation framework from the ground up; owns its architecture, reusable utilities, maintenance and ongoing evolution.',
'Automates UI and REST API workflows with Playwright, Python Requests and pytest-bdd, validating authentication, end-to-end data mapping and UASM activation through reusable BDD scenarios.',
'Architects Bitbucket CI/CD workflows for automatic UI/API regression, execution reporting, release-quality gates and repeatable release validation.',
'Develops API performance and load tests with k6, JMeter and JavaScript, and browser-workflow performance tests with JMeter, Selenium and Groovy.',
'Investigates failures and validates service/data behavior using AWS CloudWatch, DynamoDB, S3, Lambda and Splunk, connecting automation evidence with cloud logs and test data.',
'Establishes a specification-driven quality gate before implementation: reviews acceptance criteria, edge cases, negative paths, dependencies, integration behavior, test data and testability to reduce ambiguity and downstream rework.',
'Integrates local LLMs for AI-assisted test design, edge-case discovery and failure/log analysis; generates PII-free synthetic test data and reviews model outputs before use.',
'Orchestrates n8n workflows connecting Jira and QA/testing tools; coordinates automation dependencies, regression readiness and release validation across multiple teams.'
]
for b in exp:bullet(b)
labelled('Recognition','Spot Award(s) and Best Employee recognition at Experian.','Skill')

heading('PUBLIC AUTOMATION SAMPLES')
projects=[('Playwright UI/API framework','Playwright-Pytest-BDD-Framework','Playwright, pytest-bdd, Requests, page/client layers and CI gates.'),('Backend performance','Backend-Performance-Framework','k6/JMeter suites, data correlation and result gates.'),('Browser performance','UI-Performance-Framework','JMeter, Selenium and Groovy journey timing and quality gates.')]
for title,path,desc in projects:
 para=p('','Skill');link(para,title,args.verified_portfolio.rstrip('/')+'/#projects' if args.verified_portfolio else GH);para.add_run(' - '+desc)

heading('PROFESSIONAL EXPERIENCE - CONTINUED')
doc.paragraphs[-1].paragraph_format.page_break_before=True
role('Automation Test Engineer','Capgemini Engineering | Nokia WING','Dec 2021 - Jul 2025')
cap=[
'Led Selenium (Python) and Robot Framework automation across a telecom regression scope exceeding 50,000 test cases; achieved 95% automation coverage, reduced regression time by 60% and saved 300+ hours of manual effort per month.',
'Built reusable keyword-driven Robot Framework libraries, reducing automation redundancy by 25%; developed 100+ Karate REST/GraphQL API tests, reducing scripting effort by 40%.',
'Identified 10+ vulnerabilities with OWASP ZAP and validated remediation through API automation; contributed to a reported 30% improvement in defect detection.',
'Introduced Playwright, reducing cross-browser execution time by 40% compared with the previous approach.',
'Integrated suites into Jenkins and Docker-based CI/CD, reducing pipeline runtime by 25%; introduced Allure reporting and mentored 3 junior engineers.'
]
for b in cap:bullet(b)
heading('TECHNICAL SKILLS')
skills=[
('Automation & Programming','Python, Playwright, Selenium WebDriver, pytest, pytest-bdd, Robot Framework, JavaScript, Groovy, Cypress; Page Object Model, BDD, data-driven testing, parallel execution'),
('API & Integration Testing','Python Requests, Karate DSL, REST API, GraphQL, Postman, Swagger, API mocking, authentication/authorization, JSON/XML and data-mapping validation'),
('Performance Engineering','k6, JMeter, Selenium/Groovy UI performance testing, Locust, LoadRunner'),
('Security Testing','OWASP ZAP, Burp Suite, Defensics; vulnerability verification'),
('CI/CD & DevOps','Pipeline architecture, Bitbucket, Jenkins, GitHub Actions, Git, Docker, Kubernetes; regression quality gates and release validation'),
('AWS & Observability','CloudWatch, DynamoDB, S3, Lambda, Splunk; test observability and log analysis'),
('AI / LLM / Workflow Automation','n8n, local LLM integration, AI-assisted test design, prompt engineering for QA, PII-free synthetic test data, LLM-assisted edge-case discovery and defect investigation'),
('Quality Engineering Practices','Test strategy, risk-based testing, specification-driven development, shift-left QA, requirement review, acceptance-criteria refinement, regression, integration and end-to-end testing, Agile, Jira, TestRail'),
('Databases & Reporting','SQL, MySQL, PostgreSQL, MongoDB, Allure, Extent Reports')]
for label,text in skills:labelled(label,text)
heading('EDUCATION & CERTIFICATIONS')
labelled('B.Tech, Information Technology','QIS College of Engineering and Technology, Ongole | 2016 - 2020')
p('Python Programming - Kosmik Technologies (2021) | Google IT Automation with Python - Coursera (2022) | Jenkins CI/CD Pipeline - Coursera (2021)','Skill')
labelled('Academic project','Movie Suggesting System using Python and entity-based filtering.')
f=sec.footer.paragraphs[0];f.style=doc.styles['Footer'];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
f.add_run('Eswar Sai Kiran Singamsetty | ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
doc.core_properties.author='Eswar Sai Kiran Singamsetty';doc.core_properties.title='Senior QA Engineer | SDET | Automation & Quality Engineering'
doc.core_properties.subject='QA and SDET master resume';doc.core_properties.keywords='Playwright, Python, API automation, pytest-bdd, CI/CD, quality engineering'
outdir=Path(args.output_dir);outdir.mkdir(parents=True,exist_ok=True)
out=outdir/'Eswar_Sai_Kiran_Singamsetty_Master_Resume.docx';doc.save(out)
(outdir/'resume-content.json').write_text(json.dumps({'experian':exp,'capgemini':cap,'skills':skills,'projects':projects,'github':GH,'linkedin':LI},indent=2))
print(out)
